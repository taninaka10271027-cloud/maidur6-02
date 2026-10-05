# -*- coding: utf-8 -*-
"""停電作業計画書 用 シートビルダー（A4印刷に合わせて行高・改ページを自動計算）"""
import math, unicodedata
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.pagebreak import Break, RowBreak

MIN = 'ＭＳ 明朝'
GOT = 'ＭＳ ゴシック'
BASE = 10.5
F_HEAD = PatternFill('solid', fgColor='DCE6F1')
F_LABEL = PatternFill('solid', fgColor='F2F2F2')
F_NOTE = PatternFill('solid', fgColor='FFF2CC')
TH = Side(style='thin', color='000000')
HA = Side(style='hair', color='000000')
BOX = Border(left=TH, right=TH, top=TH, bottom=TH)

# 印刷可能高さ(pt)：A4縦 上下余白0.5in・ヘッダ/フッタ込み、A4横
PAGE_H = {'portrait': 752.0, 'landscape': 508.0}
# 印刷可能幅(px@96dpi) 左右余白0.4in
PAGE_W = {'portrait': 690, 'landscape': 1000}
import os
FIT = False  # 倍率100%固定（列幅はA4印刷幅内に設計済み）


def col_px(w):
    """Excel列幅(文字数)→ピクセル（既定フォント最大桁幅7px）"""
    return int(w * 7 + 5)


def text_px(s, size):
    fw = size * 96 / 72.0 * 1.06
    t = 0.0
    for ch in s:
        if unicodedata.east_asian_width(ch) in ('F', 'W', 'A'):
            t += fw
        else:
            t += fw * 0.55
    return t


def n_lines(text, width_px, size):
    if text is None or text == '':
        return 1
    avail = max(width_px - 10, 20)
    total = 0
    for para in str(text).split('\n'):
        w = text_px(para, size)
        total += max(1, math.ceil(w / avail - 0.02))
    return total


def line_h(size):
    return size * 1.36


class Sheet:
    def __init__(self, wb, title, widths, orient='portrait', first=False, footer_title=None):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = title
        self.widths = widths
        self.ncol = len(widths)
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[CL(i)].width = w
        self.orient = orient
        self.page_h = PAGE_H[orient]
        self.y = 0.0
        self.r = 1
        self.breaks = []
        total = sum(col_px(w) for w in widths)
        assert total <= PAGE_W[orient], (title, total, PAGE_W[orient])
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = orient
        if FIT:
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
        else:
            ws.page_setup.scale = 100
        ws.print_options.horizontalCentered = True
        m = ws.page_margins
        m.left = m.right = 0.4
        m.top = 0.5
        m.bottom = 0.5
        m.header = 0.25
        m.footer = 0.25
        ws.oddFooter.center.text = '&P / &N'
        ws.oddFooter.center.size = 9
        ws.oddFooter.center.font = MIN
        if footer_title:
            ws.oddHeader.right.text = footer_title
            ws.oddHeader.right.size = 8
            ws.oddHeader.right.font = MIN

    # ---------- 低レベル ----------
    def span_px(self, c0, span):
        return sum(col_px(self.widths[i]) for i in range(c0 - 1, c0 - 1 + span))

    def new_page(self):
        if self.r > 1:
            self.breaks.append(self.r - 1)
        self.y = 0.0

    def need(self, h):
        if self.y + h > self.page_h and self.y > 0:
            self.new_page()
            return True
        return False

    def put(self, cells, height=None, min_h=None, keep=False):
        """cells: list of dict(t=text, c=col(1-based), s=span, ...style)"""
        h = 0
        for d in cells:
            size = d.get('size', BASE)
            if d.get('t') is None:
                continue
            lines = n_lines(d['t'], self.span_px(d['c'], d.get('s', 1)), size)
            h = max(h, lines * line_h(size) + 5)
        if height:
            h = height
        h = max(h, min_h or 18)
        if not keep:
            self.need(h)
        r = self.r
        ws = self.ws
        for d in cells:
            c = d['c']
            s = d.get('s', 1)
            cell = ws.cell(row=r, column=c)
            cell.value = d.get('t')
            cell.font = Font(name=d.get('font', MIN), size=d.get('size', BASE), bold=d.get('b', False),
                             color=d.get('color'))
            cell.alignment = Alignment(horizontal=d.get('h', 'left'), vertical=d.get('v', 'center'),
                                       wrap_text=True, indent=d.get('ind', 0))
            if d.get('fill'):
                for cc in range(c, c + s):
                    ws.cell(row=r, column=cc).fill = d['fill']
            if d.get('border', False):
                for cc in range(c, c + s):
                    ws.cell(row=r, column=cc).border = BOX
            if s > 1:
                ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + s - 1)
        ws.row_dimensions[r].height = round(h, 2)
        self.y += h
        self.r += 1
        return r

    # ---------- 高レベル ----------
    def gap(self, h=8):
        if self.y == 0:
            return
        self.ws.row_dimensions[self.r].height = h
        self.y += h
        self.r += 1

    def title(self, text, size=14):
        self.put([dict(t=text, c=1, s=self.ncol, font=GOT, size=size, b=True)], min_h=26)

    def section(self, text, size=12, follow=60):
        # 見出しだけがページ末尾に残らないよう、後続分の高さも確保
        self.need(24 + follow)
        self.put([dict(t=text, c=1, s=self.ncol, font=GOT, size=size, b=True)], min_h=22)

    def para(self, text, c=2, size=BASE, **kw):
        return self.put([dict(t=text, c=c, s=self.ncol - c + 1, size=size, v='top', **kw)])

    def paras(self, items, c=2, **kw):
        for t in items:
            self.para(t, c=c, **kw)

    def kv(self, rows, kcol=1, kspan=1, size=BASE):
        """左:項目 右:内容 の表（罫線付）"""
        for k, v in rows:
            self.put([dict(t=k, c=kcol, s=kspan, border=True, fill=F_LABEL, h='center', size=size),
                      dict(t=v, c=kcol + kspan, s=self.ncol - kcol - kspan + 1, border=True, size=size)])

    def table(self, header, rows, spans, start=1, aligns=None, size=BASE, hsize=None, min_h=20,
              row_fills=None, header_h=None):
        """spans: 各列の結合数。header/rows: 文字列リスト"""
        hsize = hsize or size
        cols = []
        c = start
        for s in spans:
            cols.append((c, s))
            c += s
        aligns = aligns or ['left'] * len(spans)

        def hdr():
            cells = [dict(t=t, c=cc, s=ss, border=True, fill=F_HEAD, h='center', font=GOT, b=True, size=hsize)
                     for t, (cc, ss) in zip(header, cols)]
            return cells

        if header:
            # ヘッダ＋1行目は同じページへ
            first_h = 0
            if rows:
                first_h = self._row_h(rows[0], cols, size, min_h)
            self.need(30 + first_h)
            self.put(hdr(), height=header_h, keep=True)
        for i, row in enumerate(rows):
            h = self._row_h(row, cols, size, min_h)
            if self.need(h) and header:
                self.put(hdr(), height=header_h, keep=True)
            cells = []
            for j, (t, (cc, ss)) in enumerate(zip(row, cols)):
                d = dict(t=t, c=cc, s=ss, border=True, h=aligns[j], size=size)
                if row_fills and row_fills[i] and row_fills[i][j]:
                    d['fill'] = row_fills[i][j]
                cells.append(d)
            self.put(cells, height=h, keep=True)

    def _row_h(self, row, cols, size, min_h):
        h = 0
        for t, (cc, ss) in zip(row, cols):
            if t is None:
                continue
            h = max(h, n_lines(t, self.span_px(cc, ss), size) * line_h(size) + 4)
        return max(h, min_h)

    def finish(self, last_col=None):
        ws = self.ws
        last_row = self.r - 1
        ws.print_area = f'A1:{CL(last_col or self.ncol)}{last_row}'
        print(f'{ws.title:16s} rows={last_row} pages={len(self.breaks)+1} last_y={self.y:.0f}/{self.page_h:.0f}')
        ws.row_breaks = RowBreak()
        for b in self.breaks:
            ws.row_breaks.append(Break(id=b))
