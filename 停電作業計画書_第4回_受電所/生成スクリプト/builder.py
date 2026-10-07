"""停電作業計画書用 シートビルダー（テンプレート書式の再現・A4縦自動改ページ）"""
import math
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.properties import PageSetupProperties

MIN = 'ＭＳ 明朝'
GOT = 'ＭＳ ゴシック'
thin = Side(style='thin')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {'h': PatternFill('solid', fgColor='D9E1F2'), 'l': PatternFill('solid', fgColor='F2F2F2')}

PAGE_H = 758.0   # 印刷可能高さ(pt)の目安（A4縦・上0.5in/下0.6in・余裕込み）


def tw(s):
    n = 0
    for ch in s:
        o = ord(ch)
        n += 1 if (o < 0x2000 or 0xFF61 <= o <= 0xFF9F) else 2
    return n


class Cell:
    def __init__(self, span, text='', font='m', size=9.5, bold=False, fill=None,
                 align='l', border=True, color=None, valign='center'):
        self.span, self.text, self.font, self.size, self.bold = span, text, font, size, bold
        self.fill, self.align, self.border, self.color, self.valign = fill, align, border, color, valign


class Row:
    def __init__(self, cells, kind='x', minh=18.0, h=None, keepnext=0, header=None, image=None):
        self.cells, self.kind, self.minh, self.h = cells, kind, minh, h
        self.keepnext, self.header, self.image = keepnext, header, image


class SB:
    def __init__(self, wb, title, widths):
        self.ws = wb.create_sheet(title)
        self.widths = widths
        self.n = len(widths)
        self.rows = []

    # ---- 高さ計算 ----
    def _w(self, c0, span):
        return sum(self.widths[c0:c0 + span])

    def calc(self, row):
        if row.h is not None:
            return row.h
        h = row.minh
        c0 = 0
        for c in row.cells:
            if c.text:
                W = self._w(c0, c.span)
                cap = max(1.0, (W * 6.2 - 12) / (c.size * 1.36 / 2.0))  # 半角文字数
                lines = 0
                for seg in str(c.text).split('\n'):
                    lines += max(1, math.ceil(tw(seg) / cap))
                hh = lines * c.size * 1.33 + 6 + (c.size * 0.4 if lines >= 3 else 0)
                h = max(h, hh)
            c0 += c.span
        return round(h, 1)

    # ---- 部品 ----
    def add(self, row):
        self.rows.append(row)
        return row

    def blank(self, h=8):
        self.add(Row([Cell(self.n, '', border=False)], kind='blank', h=h))

    def title(self, text):
        self.add(Row([Cell(self.n, text, 'g', 14, True, border=False)], kind='title', minh=26, keepnext=2))

    def sec(self, text, size=10.5):
        self.add(Row([Cell(self.n, text, 'g', size, True, border=False)], kind='sec', minh=20, keepnext=4))

    def para(self, text, indent=1, size=9.5, keepnext=0):
        cells = []
        if indent:
            cells.append(Cell(indent, '', border=False))
        cells.append(Cell(self.n - indent, text, 'm', size, border=False))
        self.add(Row(cells, kind='para', minh=17, keepnext=keepnext))

    def note(self, text, size=9.0):
        self.add(Row([Cell(self.n, text, 'm', size, border=False)], kind='para', minh=16))

    def kv(self, label, text, lspan, size=9.5):
        self.add(Row([Cell(lspan, label, 'g', size, border=False),
                      Cell(self.n - lspan, text, 'm', size, border=False)], kind='para', minh=19))

    def table(self, spans, header, rows, size=9.0, hsize=None, aligns=None, minh=20, hminh=20,
              first_fill=None, repeat=True):
        hsize = hsize or size
        hdr = None
        if header:
            hdr = Row([Cell(s, t, 'g', hsize, fill='h', align='c') for s, t in zip(spans, header)],
                      kind="th", minh=hminh, keepnext=2)
            self.add(hdr)
        for r in rows:
            cells = []
            for i, (s, t) in enumerate(zip(spans, r)):
                al = aligns[i] if aligns else 'l'
                fill = first_fill if (i == 0 and first_fill) else None
                cells.append(Cell(s, t, 'm', size, fill=fill, align=al))
            self.add(Row(cells, kind='td', minh=minh, header=hdr if repeat else None))
        return hdr

    def image(self, path, width_px, caption=None):
        from PIL import Image
        im = Image.open(path)
        h_px = im.height * width_px / im.width
        h_pt = h_px * 0.75
        # 409pt を超えないよう複数行に分割（1ブロック）
        k = max(1, math.ceil(h_pt / 400))
        rows = []
        for i in range(k):
            rows.append(Row([Cell(self.n, '', border=False)], kind='img', h=round(h_pt / k + 0.5, 1)))
        rows[0].image = (path, width_px, h_px)
        rows[0].block = rows
        for r in rows:
            self.add(r)
        return rows

    # ---- 改ページ・出力 ----
    def paginate(self):
        out = []
        breaks = []
        cur = 0.0
        i = 0
        R = self.rows
        for r in R:
            r.hh = self.calc(r)
        while i < len(R):
            r = R[i]
            if r.kind == 'img' and getattr(r, 'block', None):
                need = sum(x.hh for x in r.block)
                grp = r.block
            else:
                need = r.hh
                grp = [r]
                j = i
                k = r.keepnext
                while k > 0 and j + 1 < len(R):
                    j += 1
                    need += R[j].hh
                    k -= 1
                    k = max(k, R[j].keepnext)
            if cur + need > PAGE_H and cur > 0:
                breaks.append(len(out))
                cur = 0.0
                # 先頭が空白行ならスキップ
                if r.kind == 'blank':
                    i += 1
                    continue
                if r.kind == 'td' and r.header is not None:
                    hc = Row(r.header.cells, kind='th', minh=r.header.minh)
                    hc.hh = r.header.hh
                    out.append(hc)
                    cur += hc.hh
            for g in grp:
                out.append(g)
                cur += g.hh
            i += len(grp)
        return out, breaks

    def build(self, print_title=None):
        ws = self.ws
        from openpyxl.utils import get_column_letter as L
        for i, w in enumerate(self.widths):
            ws.column_dimensions[L(i + 1)].width = w
        out, breaks = self.paginate()
        for ri, row in enumerate(out, start=1):
            ws.row_dimensions[ri].height = row.hh
            c0 = 1
            for c in row.cells:
                a = ws.cell(ri, c0)
                a.value = c.text if c.text != '' else None
                f = Font(name=GOT if c.font == 'g' else MIN, size=c.size, bold=c.bold,
                         color=c.color)
                al = Alignment(horizontal={'l': 'left', 'c': 'center', 'r': 'right'}[c.align],
                               vertical=c.valign, wrap_text=True)
                for cc in range(c0, c0 + c.span):
                    x = ws.cell(ri, cc)
                    x.font = f
                    x.alignment = al
                    if c.border:
                        x.border = BOX
                    if c.fill:
                        x.fill = FILL[c.fill]
                if c.span > 1:
                    ws.merge_cells(start_row=ri, start_column=c0, end_row=ri, end_column=c0 + c.span - 1)
                c0 += c.span
            if row.image:
                path, wpx, hpx = row.image
                img = XLImage(path)
                img.width, img.height = wpx, hpx
                ws.add_image(img, f'A{ri}')
        for b in breaks:
            if b > 0:
                ws.row_breaks.append(Break(id=b))
        last = len(out)
        ws.print_area = f'A1:{L(self.n)}{last}'
        ps = ws.page_setup
        ps.paperSize = 9
        ps.orientation = 'portrait'
        ps.fitToWidth = 1
        ps.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.6
        ws.page_margins.header = 0.2
        ws.page_margins.footer = 0.25
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = '- &P / &N -'
        ws.oddFooter.center.size = 9
        ws.sheet_view.view = 'pageBreakPreview'
        ws.sheet_view.zoomScale = 85
        return last
